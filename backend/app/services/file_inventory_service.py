"""Inventario y verificación de persistencia de archivos del sistema (WP-05).

Recorre las cuatro fuentes de archivos gestionados y verifica para cada uno:
existencia en disco, coincidencia de tamaño con la BD y coincidencia de hash
SHA-256 (recalculado solo si el tamaño coincide, para no leer en vano).
Las rutas devueltas al cliente son siempre lógicas (carpeta + nombre), nunca
absolutas.
"""

import io
from pathlib import Path

import pandas as pd
from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader

from app.core.config import settings
from app.models.document import DocumentVersion
from app.models.dynamic_field import FieldAttachment, TemplateVersion
from app.services.template_docx_service import sha256_hex

MAX_HASH_BYTES = 10 * 1024 * 1024
MAX_PREVIEW_WORDS = 4000
PREVIEWABLE_SUFFIXES = {".docx", ".pdf", ".csv", ".xlsx"}

_KIND_SPECS = {
    "TEMPLATE_VERSION": {
        "logical_dir": "uploads/templates/",
        "media_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    },
    "ATTACHMENT": {"logical_dir": "uploads/attachments/", "media_type": None},
    "DOCUMENT_VERSION": {
        "logical_dir": "generated/documents/",
        "media_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    },
    "PREVIEW": {"logical_dir": "generated/previews/", "media_type": None},
}


def _verify_file(path: Path, db_size: int | None, db_hash: str | None) -> dict:
    """Verificación de persistencia de un archivo contra su registro."""
    exists = path.is_file()
    size_on_disk = path.stat().st_size if exists else None
    size_matches = (
        exists and db_size is not None and size_on_disk == db_size
        if db_size is not None
        else exists
    )
    hash_matches = None
    if exists and db_hash and size_matches and (size_on_disk or 0) <= MAX_HASH_BYTES:
        hash_matches = sha256_hex(path.read_bytes()) == db_hash

    if not exists:
        status = "NO_ENCONTRADO"
    elif hash_matches is False or (db_size is not None and not size_matches):
        status = "DIFIERE"
    else:
        status = "OK"
    return {
        "exists_on_disk": exists,
        "size_on_disk": size_on_disk,
        "size_matches_db": bool(size_matches),
        "hash_matches_db": hash_matches,
        "status": status,
    }


def _logical_path(kind: str, stored_path: str | None) -> str:
    """Carpeta lógica + nombre (jamás ruta absoluta del servidor)."""
    name = Path(stored_path).name if stored_path else "—"
    return f"{_KIND_SPECS[kind]['logical_dir']}{name}"


def build_inventory(db, kind: str | None, status_filter: str | None) -> list[dict]:
    """Inventario completo de archivos gestionados con su verificación."""
    rows: list[dict] = []

    if kind in (None, "TEMPLATE_VERSION"):
        versions = (
            db.query(TemplateVersion)
            .filter(TemplateVersion.file_path.isnot(None))
            .order_by(TemplateVersion.created_at.desc())
            .all()
        )
        for version in versions:
            path = Path(version.file_path)
            verification = _verify_file(path, version.file_size, version.file_hash)
            rows.append(
                {
                    "kind": "TEMPLATE_VERSION",
                    "record_id": version.id,
                    "file_name": version.original_filename or path.name,
                    "reference": f"Plantilla v{version.version_number} ({version.status})",
                    "logical_path": _logical_path("TEMPLATE_VERSION", version.file_path),
                    "db_size": version.file_size,
                    "db_hash": version.file_hash,
                    "previewable": path.suffix.lower() in PREVIEWABLE_SUFFIXES,
                    "created_at": version.created_at.isoformat(),
                    **verification,
                }
            )

    if kind in (None, "ATTACHMENT"):
        attachments = (
            db.query(FieldAttachment)
            .order_by(FieldAttachment.created_at.desc())
            .all()
        )
        for attachment in attachments:
            path = settings.UPLOAD_DIR / "attachments" / attachment.storage_name
            verification = _verify_file(path, attachment.size, None)
            rows.append(
                {
                    "kind": "ATTACHMENT",
                    "record_id": attachment.id,
                    "file_name": attachment.original_name,
                    "reference": f"Adjunto del campo {attachment.field_key}",
                    "logical_path": _logical_path("ATTACHMENT", attachment.storage_name),
                    "db_size": attachment.size,
                    "db_hash": None,
                    "previewable": Path(attachment.original_name).suffix.lower()
                    in PREVIEWABLE_SUFFIXES,
                    "created_at": attachment.created_at.isoformat(),
                    **verification,
                }
            )

    if kind in (None, "DOCUMENT_VERSION"):
        documents = (
            db.query(DocumentVersion)
            .order_by(DocumentVersion.created_at.desc())
            .all()
        )
        for document in documents:
            path = Path(document.file_path)
            verification = _verify_file(path, document.file_size, document.file_hash)
            rows.append(
                {
                    "kind": "DOCUMENT_VERSION",
                    "record_id": document.id,
                    "file_name": f"borrador_v{document.version_number}.docx",
                    "reference": f"Borrador v{document.version_number} ({document.validation_status})",
                    "logical_path": _logical_path("DOCUMENT_VERSION", document.file_path),
                    "db_size": document.file_size,
                    "db_hash": document.file_hash,
                    "previewable": True,
                    "created_at": document.created_at.isoformat(),
                    **verification,
                }
            )

    if kind in (None, "PREVIEW"):
        previews_dir = settings.GENERATED_DIR / "previews"
        if previews_dir.is_dir():
            for file in sorted(
                previews_dir.glob("*.docx"),
                key=lambda f: f.stat().st_mtime,
                reverse=True,
            ):
                verification = _verify_file(file, None, None)
                rows.append(
                    {
                        "kind": "PREVIEW",
                        "record_id": file.stem,
                        "file_name": file.name,
                        "reference": "Render de prueba de plantilla",
                        "logical_path": _logical_path("PREVIEW", file.name),
                        "db_size": None,
                        "db_hash": None,
                        "previewable": True,
                        "created_at": "",
                        **verification,
                    }
                )

    if status_filter:
        rows = [row for row in rows if row["status"] == status_filter]
    return rows


# ---------------------------------------------------------------------------
# Previsualización (sin descarga)
# ---------------------------------------------------------------------------


def _resolve_record_path(db, kind: str, record_id: str) -> tuple[Path, str]:
    """Ruta real del archivo solicitado para preview, validando el tipo."""
    if kind == "TEMPLATE_VERSION":
        record = db.get(TemplateVersion, record_id)
        stored = record.file_path if record else None
        name = record.original_filename if record else None
    elif kind == "ATTACHMENT":
        record = db.get(FieldAttachment, record_id)
        stored = (
            str(settings.UPLOAD_DIR / "attachments" / record.storage_name)
            if record
            else None
        )
        name = record.original_name if record else None
    elif kind == "DOCUMENT_VERSION":
        record = db.get(DocumentVersion, record_id)
        stored = record.file_path if record else None
        name = f"borrador_v{record.version_number}.docx" if record else None
    elif kind == "PREVIEW":
        # El record_id de previews es el stem del nombre de archivo (UUID).
        if not record_id.replace("-", "").isalnum() or "/" in record_id:
            record, stored, name = None, None, None
        else:
            candidate = settings.GENERATED_DIR / "previews" / f"{record_id}.docx"
            stored = str(candidate) if candidate.is_file() else None
            name = f"{record_id}.docx"
            record = object() if stored else None
    else:
        record, stored, name = None, None, None

    if record is None or not stored:
        raise HTTPException(404, "Archivo no encontrado en el inventario.")
    path = Path(stored)
    if not path.is_file():
        raise HTTPException(404, "El archivo no existe en disco.")
    return path, name or path.name


def preview_file(db, kind: str, record_id: str) -> dict:
    """Extrae contenido textual/tabulado para mostrar en modal (sin descarga)."""
    path, name = _resolve_record_path(db, kind, record_id)
    suffix = path.suffix.lower()

    if suffix == ".docx":
        document = Document(str(path))
        parts: list[str] = []
        words = 0
        truncated = False

        def feed(text: str) -> bool:
            nonlocal words, truncated
            count = len(text.split())
            if words + count > MAX_PREVIEW_WORDS:
                truncated = True
                return False
            words += count
            return True

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                if not feed(paragraph.text):
                    break
                parts.append(paragraph.text)
        if not truncated:
            for table in document.tables:
                for row in table.rows:
                    line = " | ".join(cell.text.strip() for cell in row.cells)
                    if line.strip(" |"):
                        if not feed(line):
                            break
                        parts.append(line)

        return {
            "file_name": name,
            "kind": kind,
            "media_type": _KIND_SPECS[kind]["media_type"] or "application/octet-stream",
            "preview_type": "text",
            "content": "\n\n".join(parts),
            "truncated": truncated,
        }

    if suffix == ".pdf":
        try:
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise ValueError("encrypted")
            text = "\n\n".join(
                (page.extract_text() or "").strip()
                for page in reader.pages
                if (page.extract_text() or "").strip()
            )
        except Exception:
            text = ""
        if not text.strip():
            return {
                "file_name": name,
                "kind": kind,
                "media_type": "application/pdf",
                "preview_type": "unavailable",
                "content": "EXTRACCION_NO_DISPONIBLE_SIN_OCR",
                "truncated": False,
            }
        words = text.split()
        truncated = len(words) > MAX_PREVIEW_WORDS
        return {
            "file_name": name,
            "kind": kind,
            "media_type": "application/pdf",
            "preview_type": "text",
            "content": " ".join(words[:MAX_PREVIEW_WORDS]),
            "truncated": truncated,
        }

    if suffix == ".csv":
        frame = pd.read_csv(path).head(25)
        return {
            "file_name": name,
            "kind": kind,
            "media_type": "text/csv",
            "preview_type": "table",
            "content": {
                "columns": list(frame.columns),
                "rows": frame.astype(str).values.tolist(),
            },
            "truncated": False,
        }

    if suffix == ".xlsx":
        frame = pd.read_excel(path, engine="openpyxl").head(25)
        return {
            "file_name": name,
            "kind": kind,
            "media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "preview_type": "table",
            "content": {
                "columns": list(frame.columns),
                "rows": frame.astype(str).values.tolist(),
            },
            "truncated": False,
        }

    raise HTTPException(
        415, f"Previsualización no disponible para archivos '{suffix}'."
    )
