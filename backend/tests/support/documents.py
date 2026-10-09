"""Constructores de DOCX sintéticos con párrafos y tablas."""

from io import BytesIO
from pathlib import Path

from docx import Document


def docx_bytes(paragraphs: list[str], table_text: str | None = None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_text is not None:
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = table_text
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def write_docx(
    target: Path, paragraphs: list[str], table_text: str | None = None
) -> Path:
    target.write_bytes(docx_bytes(paragraphs, table_text))
    return target
