"""Integration tests for the system file inventory and preview (WP-05).

Verificación de persistencia (existencia, tamaño, hash) y previsualización
en modal de plantillas, adjuntos, borradores y renders de prueba.
"""

import io
from pathlib import Path

from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.security import get_password_hash
from app.models.document import Document, DocumentVersion
from app.models.dynamic_field import FieldAttachment
from app.models.user import User
from app.services.template_docx_service import sha256_hex

BASE = "/api/v1"


def _create_user(db_session: Session, suffix: str, role: str) -> User:
    user = User(
        username=f"user_{suffix}",
        email=f"user_{suffix}@bufetenotarial.demo",
        full_name=f"Usuario {suffix}",
        password_hash=get_password_hash("Admin1234!"),
        role=role,
        status="ACTIVE",
    )
    db_session.add(user)
    db_session.commit()
    return user


def _token(client: TestClient, username: str) -> str:
    resp = client.post(
        f"{BASE}/auth/login",
        json={"username_or_email": username, "password": "Admin1234!"},
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["access_token"]


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _docx_bytes(text: str) -> bytes:
    document = DocxDocument()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _seed_files(db_session: Session, tmp_upload: Path):
    """Plantilla en uploads/templates y adjunto en uploads/attachments."""
    from app.models.dynamic_field import Template, TemplateVersion

    content = _docx_bytes("ESCRITURA DE PRUEBA {{ numero }}")
    template = Template(name="Plantilla Inventario", case_type="COMPRAVENTA", status="ACTIVE")
    db_session.add(template)
    db_session.flush()
    path = tmp_upload / "templates" / "tpl_inventario.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    version = TemplateVersion(
        template_id=template.id,
        version_number=1,
        status="ACTIVA",
        file_path=str(path),
        original_filename="tpl_inventario.docx",
        file_hash=sha256_hex(content),
        file_size=len(content),
    )
    db_session.add(version)
    db_session.flush()

    attachment_content = _docx_bytes("ADJUNTO DE SOPORTE")
    attachment_path = tmp_upload / "attachments" / "adj_inv.docx"
    attachment_path.parent.mkdir(parents=True, exist_ok=True)
    attachment_path.write_bytes(attachment_content)
    attachment = FieldAttachment(
        case_id="case-x",
        template_version_id=version.id,
        field_key="soporte",
        original_name="adj_inv.docx",
        storage_name="adj_inv.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        size=len(attachment_content),
    )
    db_session.add(attachment)

    doc_row = Document(case_id="case-x", title="Borrador Inventario")
    db_session.add(doc_row)
    db_session.flush()
    gen_path = tmp_upload.parent / "generated" / "documents" / "gen_inv.docx"
    gen_path.parent.mkdir(parents=True, exist_ok=True)
    gen_content = _docx_bytes("BORRADOR GENERADO")
    gen_path.write_bytes(gen_content)
    doc_version = DocumentVersion(
        document_id=doc_row.id,
        version_number=1,
        template_version_id=version.id,
        data_snapshot={},
        file_path=str(gen_path),
        file_hash=sha256_hex(gen_content),
        file_size=len(gen_content),
        created_by_id="user-x",
    )
    db_session.add(doc_version)
    db_session.commit()
    return version, attachment, doc_version


def test_inventory_lists_sources_with_verification(
    client: TestClient, db_session: Session, tmp_path: Path, monkeypatch
):
    _create_user(db_session, "files1", "ADMINISTRADOR")
    token = _token(client, "user_files1")
    monkeypatch.setattr(settings, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(settings, "GENERATED_DIR", tmp_path / "generated")
    _seed_files(db_session, tmp_path / "uploads")

    resp = client.get(f"{BASE}/files/inventory", headers=_auth(token))
    assert resp.status_code == 200, resp.text
    items = resp.json()["items"]
    kinds = {item["kind"] for item in items}
    assert {"TEMPLATE_VERSION", "ATTACHMENT", "DOCUMENT_VERSION"} <= kinds
    assert all(item["status"] == "OK" for item in items)
    assert all(item["exists_on_disk"] for item in items)
    template_row = next(i for i in items if i["kind"] == "TEMPLATE_VERSION")
    assert template_row["hash_matches_db"] is True
    assert not str(template_row["logical_path"]).startswith("C:")  # jamás ruta absoluta


def test_inventory_marks_missing_file(client: TestClient, db_session: Session, tmp_path: Path, monkeypatch):
    _create_user(db_session, "files2", "ADMINISTRADOR")
    token = _token(client, "user_files2")
    monkeypatch.setattr(settings, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(settings, "GENERATED_DIR", tmp_path / "generated")
    version, _, _ = _seed_files(db_session, tmp_path / "uploads")

    # Borrar el archivo de la plantilla: el registro debe marcar NO_ENCONTRADO.
    Path(version.file_path).unlink()
    resp = client.get(f"{BASE}/files/inventory?kind=TEMPLATE_VERSION", headers=_auth(token))
    assert resp.status_code == 200
    row = resp.json()["items"][0]
    assert row["status"] == "NO_ENCONTRADO"
    assert row["exists_on_disk"] is False


def test_preview_docx_returns_text(client: TestClient, db_session: Session, tmp_path: Path, monkeypatch):
    _create_user(db_session, "files3", "ADMINISTRADOR")
    token = _token(client, "user_files3")
    monkeypatch.setattr(settings, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(settings, "GENERATED_DIR", tmp_path / "generated")
    version, _, _ = _seed_files(db_session, tmp_path / "uploads")

    resp = client.get(
        f"{BASE}/files/preview/TEMPLATE_VERSION/{version.id}", headers=_auth(token)
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["preview_type"] == "text"
    assert "ESCRITURA DE PRUEBA" in data["content"]
    assert data["file_name"] == "tpl_inventario.docx"


def test_preview_document_version(client: TestClient, db_session: Session, tmp_path: Path, monkeypatch):
    _create_user(db_session, "files4", "ADMINISTRADOR")
    token = _token(client, "user_files4")
    monkeypatch.setattr(settings, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(settings, "GENERATED_DIR", tmp_path / "generated")
    _, _, doc_version = _seed_files(db_session, tmp_path / "uploads")

    resp = client.get(
        f"{BASE}/files/preview/DOCUMENT_VERSION/{doc_version.id}", headers=_auth(token)
    )
    assert resp.status_code == 200
    assert "BORRADOR GENERADO" in resp.json()["content"]


def test_preview_unknown_record_is_404(client: TestClient, db_session: Session):
    _create_user(db_session, "files5", "ADMINISTRADOR")
    token = _token(client, "user_files5")
    resp = client.get(
        f"{BASE}/files/preview/TEMPLATE_VERSION/00000000-0000-0000-0000-000000000000",
        headers=_auth(token),
    )
    assert resp.status_code == 404


def test_inventory_forbidden_for_auxiliar(client: TestClient, db_session: Session):
    _create_user(db_session, "files6", "AUXILIAR")
    token = _token(client, "user_files6")
    resp = client.get(f"{BASE}/files/inventory", headers=_auth(token))
    assert resp.status_code == 403
