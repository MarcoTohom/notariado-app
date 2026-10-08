"""Integration tests for the editor live-preview endpoint (WP-07)."""

import io

from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.core.security import get_password_hash
from app.models.user import User

BASE = "/api/v1"


def _create_user(db_session: Session, suffix: str, role: str) -> User:
    user = User(
        username=f"user_{suffix}",
        email=f"user_{suffix}@bufetenotarial.demo",
        full_name=f"Usuario {suffix}",
        password_hash=get_password_hash("Admin1234!"),
        role=role,
        status="ACTIVE",
    )
    db_session.add(user)
    db_session.commit()
    return user


def _token(client: TestClient, username: str) -> str:
    resp = client.post(
        f"{BASE}/auth/login",
        json={"username_or_email": username, "password": "Admin1234!"},
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["access_token"]


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _setup(client: TestClient, db_session: Session, suffix: str):
    _create_user(db_session, suffix, "ADMINISTRADOR")
    token = _token(client, f"user_{suffix}")

    resp = client.post(
        f"{BASE}/clients",
        json={"first_name": "Ana", "last_name": "Lopez", "dpi": "1234567890101"},
        headers=_auth(token),
    )
    client_id = resp.json()["id"]
    resp = client.post(
        f"{BASE}/cases",
        json={
            "title": "Compraventa para preview en vivo",
            "case_type": "COMPRAVENTA",
            "parties": [{"client_id": client_id, "party_role": "COMPRADOR"}],
        },
        headers=_auth(token),
    )
    case_id = resp.json()["id"]

    document = DocxDocument()
    document.add_paragraph("ESCRITURA No. {{ numero_escritura }}")
    document.add_paragraph("COMPRADOR: {{ comprador.nombre }} DPI {{ comprador.dpi }}")
    document.add_paragraph("NOTAS: {{ observaciones }}")
    buffer = io.BytesIO()
    document.save(buffer)
    resp = client.post(
        f"{BASE}/templates",
        data={"name": f"Preview {suffix}", "case_type": "COMPRAVENTA"},
        files={"file": ("preview.docx", buffer.getvalue(), "application/octet-stream")},
        headers=_auth(token),
    )
    assert resp.status_code == 201, resp.text
    version_id = resp.json()["versions"][0]["id"]
    return token, case_id, version_id


def test_preview_render_substitutes_values_as_html(
    client: TestClient, db_session: Session
):
    token, case_id, version_id = _setup(client, db_session, "prev1")
    resp = client.post(
        f"{BASE}/documents/preview-render",
        json={
            "case_id": case_id,
            "template_version_id": version_id,
            "values": {
                "numero_escritura": "151",
                "comprador_nombre": "Ana Lopez",
                "comprador_dpi": "1234567890101",
                "observaciones": "sin observaciones",
            },
        },
        headers=_auth(token),
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert "<p>ESCRITURA No. 151</p>" in data["html"]
    assert "Ana Lopez" in data["html"]
    assert "1234567890101" in data["html"]
    assert data["placeholders_free"] is True


def test_preview_render_escapes_html_injection(client: TestClient, db_session: Session):
    """El texto del usuario jamás se interpreta como HTML (anti-XSS)."""
    token, case_id, version_id = _setup(client, db_session, "prev2")
    resp = client.post(
        f"{BASE}/documents/preview-render",
        json={
            "case_id": case_id,
            "template_version_id": version_id,
            "values": {"observaciones": "<script>alert('xss')</script>"},
        },
        headers=_auth(token),
    )
    assert resp.status_code == 200
    html = resp.json()["html"]
    assert "<script>alert" not in html
    assert "&lt;script&gt;" in html


def test_preview_render_requires_auth(client: TestClient):
    resp = client.post(
        f"{BASE}/documents/preview-render",
        json={"case_id": "x", "template_version_id": "y", "values": {}},
    )
    assert resp.status_code == 401
