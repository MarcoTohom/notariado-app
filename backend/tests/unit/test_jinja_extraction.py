"""Pruebas unitarias del extractor léxico Jinja2 y utilidades de Fase 5.

Cubre el paso 2 del skill docx-template: detección de variables simples,
bucles {% for %} y condicionales {% if %} en párrafos y tablas, además de
la sugerencia de tipos, la construcción del contexto sintético y la
verificación de placeholders residuales (RULE-017).
"""

import io
from pathlib import Path

import pytest
from docx import Document

from app.services.docx.analysis import (
    JinjaExtraction,
    extract_jinja_variables,
    find_residual_variables,
    humanize_label,
    suggest_field_type,
)
from app.services.docx.context import build_sample_context
from app.services.docx.files import sanitize_original_name


def _make_docx(
    tmp_path: Path, paragraphs: list[str], table_text: str | None = None
) -> Path:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_text is not None:
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = table_text
    target = tmp_path / "plantilla_prueba.docx"
    document.save(str(target))
    return target


# ---------------------------------------------------------------------------
# Extracción de variables
# ---------------------------------------------------------------------------


def test_extract_simple_variables_in_order_and_deduped(tmp_path: Path):
    path = _make_docx(
        tmp_path,
        [
            "ESCRITURA No. {{ numero_escritura }}",
            "COMPARECE: {{ comprador.nombre_completo }}, DPI {{ comprador.dpi }}.",
            "Repite el número: {{numero_escritura}}",
        ],
    )
    extraction = extract_jinja_variables(path)
    assert extraction.variables == [
        "numero_escritura",
        "comprador.nombre_completo",
        "comprador.dpi",
    ]


def test_extract_variables_inside_tables(tmp_path: Path):
    path = _make_docx(
        tmp_path, ["Sin variables aquí"], table_text="NIT: {{ vendedor.nit }}"
    )
    extraction = extract_jinja_variables(path)
    assert extraction.variables == ["vendedor.nit"]


def test_extract_loops_and_conditionals(tmp_path: Path):
    path = _make_docx(
        tmp_path,
        [
            "{% for comp in compradores %}{{ comp.nombre }} ({{ comp.dpi }}){% endfor %}",
            "{% if comprador.nit %}NIT: {{ comprador.nit }}{% endif %}",
        ],
    )
    extraction = extract_jinja_variables(path)
    assert extraction.loops == [{"item": "comp", "collection": "compradores"}]
    assert extraction.conditionals == ["comprador.nit"]
    assert "comp.nombre" in extraction.variables
    assert "comp.dpi" in extraction.variables


def test_extract_ignores_plain_text_and_single_braces(tmp_path: Path):
    path = _make_docx(tmp_path, ["Texto {sin} variables y { otro } corchete."])
    extraction = extract_jinja_variables(path)
    assert extraction.variables == []
    assert extraction.loops == []
    assert extraction.conditionals == []


# ---------------------------------------------------------------------------
# Sugerencia de tipos y etiquetas
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("variable", "expected"),
    [
        ("comprador.dpi", "dpi"),
        ("cui_compareciente", "dpi"),
        ("empresa.nit", "nit"),
        ("cliente.correo", "email"),
        ("telefono_contacto", "phone"),
        ("fecha_escritura", "date"),
        ("precio_inmueble", "currency"),
        ("monto_total", "currency"),
        ("porcentaje_aporte", "percentage"),
        ("vendedor.nombre_completo", "name"),
        ("direccion_inmueble", "textarea"),
        ("numero_escritura", "text"),
    ],
)
def test_suggest_field_type(variable: str, expected: str):
    assert suggest_field_type(variable) == expected


def test_humanize_label_handles_paths_and_acronyms():
    assert humanize_label("comprador.dpi") == "Comprador - DPI"
    assert humanize_label("fecha_escritura") == "Fecha Escritura"
    assert humanize_label("vendedor.nit") == "Vendedor - NIT"


# ---------------------------------------------------------------------------
# Contexto sintético para render de prueba
# ---------------------------------------------------------------------------


def test_build_sample_context_scalars_and_nested():
    extraction = JinjaExtraction(
        variables=["numero_escritura", "comprador.dpi", "comprador.nombre"],
        loops=[],
        conditionals=[],
    )
    context = build_sample_context(extraction)
    assert context["numero_escritura"] == "[valor de prueba]"
    assert context["comprador"]["dpi"] == "1234567890101"
    assert context["comprador"]["nombre"] == "Nombre Sintético de Prueba"


def test_build_sample_context_loop_rows():
    extraction = JinjaExtraction(
        variables=["comp.nombre", "comp.dpi", "fecha"],
        loops=[{"item": "comp", "collection": "compradores"}],
        conditionals=["mostrar_testigos"],
    )
    context = build_sample_context(extraction)
    assert len(context["compradores"]) == 2
    assert context["compradores"][0]["dpi"] == "1234567890101"
    assert context["compradores"][1]["indice"] == 2
    assert context["mostrar_testigos"] is True


# ---------------------------------------------------------------------------
# Verificación de placeholders residuales (RULE-017)
# ---------------------------------------------------------------------------


def test_find_residual_variables_detects_leftovers(tmp_path: Path):
    path = _make_docx(tmp_path, ["Quedó sin sustituir: {{ variable_faltante }}"])
    assert find_residual_variables(path) == ["variable_faltante"]


def test_find_residual_variables_clean_document(tmp_path: Path):
    path = _make_docx(tmp_path, ["Documento completamente sustituido."])
    assert find_residual_variables(path) == []


# ---------------------------------------------------------------------------
# Sanitización de nombres de archivo (anti path-traversal)
# ---------------------------------------------------------------------------


def test_sanitize_original_name_strips_paths_and_control_chars():
    assert sanitize_original_name("../../etc/secreto.docx") == "secreto.docx"
    assert sanitize_original_name("C:\\temp\\mi plantilla.docx") == "mi plantilla.docx"
    assert (
        sanitize_original_name("nombre\x00con\x1fcontrol.docx")
        == "nombreconcontrol.docx"
    )
    assert sanitize_original_name("") == "plantilla.docx"


def test_docx_fixture_is_valid_openxml(tmp_path: Path):
    """El fixture usado en las pruebas es un DOCX abrible por python-docx."""
    path = _make_docx(tmp_path, ["Hola {{ mundo }}"])
    document = Document(str(path))
    assert document.paragraphs[0].text == "Hola {{ mundo }}"
    with open(path, "rb") as handle:
        buffer = io.BytesIO(handle.read())
    assert buffer.read(2) == b"PK"
