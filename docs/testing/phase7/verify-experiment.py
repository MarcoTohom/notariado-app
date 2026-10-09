"""Verifica corpus persistido y exportaciones sin crear mediciones ficticias."""

import io
import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "backend"))

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.core.security import create_access_token
from app.db.session import SessionLocal
from app.main import app
from app.models.experiment import TestCase, TestExecution, TimeMeasurement
from app.models.user import User
from app.rules.engine import run_rules
from app.services.validation_context import build_validation_context

with SessionLocal() as db:
    cases = db.query(TestCase).all()
    assert len(cases) == 100
    counts = Counter((case.case_type, case.has_anomalies) for case in cases)
    assert len(counts) == 10 and all(count == 10 for count in counts.values())
    for case in cases:
        context = build_validation_context(db, case.case_id, case.template_version_id, None)
        actual = {finding.rule_id for finding in run_rules(context)}
        assert actual == set(case.expected_findings), (case.case_type, actual, case.expected_findings)
    assert db.query(TestExecution).count() == 0
    assert db.query(TimeMeasurement).count() == 0
    admin = db.query(User).filter_by(role="ADMINISTRADOR").first()
    headers = {"Authorization": f"Bearer {create_access_token(admin.id)}"}

with TestClient(app) as client:
    stats_response = client.get("/api/v1/experiment/stats", headers=headers)
    assert stats_response.status_code == 200
    stats = stats_response.json()
    assert stats["system"]["n"] == stats["traditional"]["n"] == 0
    assert stats["baseline_minutes"] == 240 and stats["reduction_percentage"] is None
    xlsx = client.get("/api/v1/experiment/export.xlsx", headers=headers)
    csv = client.get("/api/v1/experiment/export.csv", headers=headers)
    assert xlsx.status_code == csv.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx.content), read_only=True)
    assert workbook.sheetnames == ["Ejecuciones", "Resumen"]
    summary = dict(list(workbook["Resumen"].values)[1:])
    assert summary["cases_total"] == 100 and summary["executions_total"] == 0
    assert summary["reduction_percentage"] is None
    assert csv.content.decode("utf-8").strip() == ""
    workbook.close()

result = dict(corpus=100, types=5, clean=50, anomalous=50, persisted_rule_results_checked=100,
              rule_results="exactly match expected findings", xlsx="passed", csv="passed (empty without executions)",
              executions=0, time_measurements=0, reduction_percentage=None, baseline_minutes=240,
              status="passed")
Path(__file__).with_name("experiment-verification.json").write_text(json.dumps(result, indent=2)+"\n", encoding="utf-8")
print(json.dumps(result, indent=2))
